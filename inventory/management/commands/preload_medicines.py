
"""
Preload Pakistan medicine catalog into MedicineMaster (NOT live stock).
Usage:
  python manage.py preload_medicines /path/to/Pakistan_Medicine_Master_List_LEGACY_prices.xlsx
  python manage.py preload_medicines --limit 5000
"""
from decimal import Decimal, InvalidOperation
from django.core.management.base import BaseCommand
from inventory.models import MedicineMaster


def money(val):
    if val is None or val == '':
        return None
    try:
        return Decimal(str(val).replace(',', '').strip()).quantize(Decimal('0.01'))
    except (InvalidOperation, ValueError):
        return None


class Command(BaseCommand):
    help = 'Import Pakistan medicine master list (catalog only, quantity=0 stock not created)'

    def add_arguments(self, parser):
        parser.add_argument('xlsx', nargs='?', default='', help='Path to xlsx file')
        parser.add_argument('--limit', type=int, default=0, help='Max rows (0=all)')
        parser.add_argument('--clear', action='store_true', help='Delete existing master rows first')

    def handle(self, *args, **opts):
        path = opts['xlsx'] or '/home/workdir/attachments/Pakistan_Medicine_Master_List_LEGACY_prices.xlsx'
        try:
            import openpyxl
        except ImportError:
            self.stderr.write('pip install openpyxl')
            return

        if opts['clear']:
            n = MedicineMaster.objects.all().delete()[0]
            self.stdout.write(f'Cleared {n} existing master rows')

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        # Prefer sheet named Medicines
        ws = wb['Medicines'] if 'Medicines' in wb.sheetnames else wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [str(h or '').strip() for h in next(rows)]
        # map columns by header keywords
        def col(*names):
            for i, h in enumerate(headers):
                hl = h.lower()
                for n in names:
                    if n.lower() in hl:
                        return i
            return None

        i_name = col('Product Name')
        i_salt = col('Ingredients', 'Salt')
        i_brand = col('Brand Name')
        i_retail = col('Legacy Retail')
        i_company = col('Company', 'Manufacturer')
        i_strength = col('Strength')
        i_form = col('Dosage Form')
        i_pack = col('Pack Size')
        i_group = col('Therapeutic')
        i_trade = col('Legacy Trade')
        i_mrp = col('Current MRP')

        if i_name is None:
            self.stderr.write(f'Product Name column not found. Headers: {headers[:15]}')
            return

        def cell(row, idx):
            if idx is None or idx >= len(row):
                return ''
            v = row[idx]
            return '' if v is None else str(v).strip()

        batch = []
        created = 0
        limit = opts['limit'] or 0
        for row in rows:
            if not row:
                continue
            pname = cell(row, i_name)
            if not pname or pname.lower() in ('product name', 's.no'):
                continue
            batch.append(MedicineMaster(
                product_name=pname[:255],
                brand_name=cell(row, i_brand)[:120],
                ingredients=cell(row, i_salt)[:500],
                company=cell(row, i_company)[:200],
                strength=cell(row, i_strength)[:120],
                dosage_form=cell(row, i_form)[:80],
                pack_size=cell(row, i_pack)[:80],
                therapeutic_group=cell(row, i_group)[:120],
                legacy_retail_price=money(cell(row, i_retail)),
                legacy_trade_price=money(cell(row, i_trade)),
                current_mrp=money(cell(row, i_mrp)),
            ))
            if len(batch) >= 1000:
                MedicineMaster.objects.bulk_create(batch, ignore_conflicts=True)
                created += len(batch)
                self.stdout.write(f'… {created}')
                batch = []
            if limit and created + len(batch) >= limit:
                break
        if batch:
            if limit:
                batch = batch[: max(0, limit - created)]
            MedicineMaster.objects.bulk_create(batch, ignore_conflicts=True)
            created += len(batch)
        self.stdout.write(self.style.SUCCESS(
            f'Preloaded {created} catalog rows into MedicineMaster. '
            f'These are NOT stock — purchase to create live Product SKUs.'
        ))
